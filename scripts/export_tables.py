"""Export all DB tables to Excel for rebate formula calculation."""
import sys, os
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from datetime import datetime
from app.db.connection import get_conn

# ── Styles ────────────────────────────────────────────────────────────────────
HDR_FILL  = PatternFill("solid", fgColor="1E3A5F")
HDR_FONT  = Font(bold=True, color="FFFFFF", size=10, name="Consolas")
TITLE_FONT= Font(bold=True, color="1E3A5F", size=12)
ALT_FILL  = PatternFill("solid", fgColor="F0F4FA")
DATA_FONT = Font(size=10, name="Consolas")
LEFT      = Alignment(horizontal="left",  vertical="center")
RIGHT_NUM = Alignment(horizontal="right", vertical="center")
CENTER    = Alignment(horizontal="center",vertical="center", wrap_text=True)
THIN      = Side(style="thin", color="BFC9D9")
BORDER    = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

# Float columns (format as #,##0.00)
FLOAT_COLS = {
    "total_volume","total_commission","rebate_amount",
    "fx_lot","fx_commission","hang_hoa_lot","hang_hoa_commission",
    "index_lot","index_commission","crypto_lot","crypto_commission",
    "sharecfd_lot","sharecfd_commission","bond_lot","bond_commission",
    "synthetic_lot","synthetic_commission",
    "total_rows","total_amount",
    "target_lot","reward_value",
}
INT_COLS = {"id","batch_id","customer_account_id","program_tier_id",
            "customer_id","broker_id","program_id","pending_program_id",
            "tier_number","display_order"}
DATE_COLS = {"period_date","created_at","uploaded_at","customer_linked_at",
             "reset_token_expires","verify_token_expires"}


def dump_table(wb, sheet_name, rows, col_names, first=False):
    ws = wb.active if first else wb.create_sheet(sheet_name)
    ws.title = sheet_name
    ws.freeze_panes = "A3"
    ws.sheet_view.showGridLines = False

    # Row 1: title
    now_str = datetime.now().strftime("%Y-%m-%d %H:%M")
    t = ws.cell(1, 1, f"{sheet_name}  |  {len(rows)} rows  |  exported {now_str}")
    t.font = TITLE_FONT
    t.alignment = LEFT

    # Row 2: headers
    for ci, col in enumerate(col_names, 1):
        c = ws.cell(2, ci, col)
        c.fill = HDR_FILL; c.font = HDR_FONT
        c.alignment = CENTER; c.border = BORDER

    ws.row_dimensions[2].height = 20

    # Data rows
    for ri, row in enumerate(rows, 3):
        alt = (ri % 2 == 0)
        for ci, col in enumerate(col_names, 1):
            raw = row[col]

            # Stringify date/datetime
            if col in DATE_COLS and raw is not None:
                val = str(raw)[:19]
            else:
                val = raw if raw is not None else (0.0 if col in FLOAT_COLS else raw)

            cell = ws.cell(ri, ci, val)
            cell.font = DATA_FONT
            cell.border = BORDER
            if alt:
                cell.fill = ALT_FILL

            if col in FLOAT_COLS:
                cell.alignment = RIGHT_NUM
                cell.number_format = "#,##0.00"
            elif col in INT_COLS:
                cell.alignment = RIGHT_NUM
                cell.number_format = "#,##0"
            else:
                cell.alignment = LEFT

    # Auto-fit column widths
    for ci, col in enumerate(col_names, 1):
        max_len = max(len(col) + 2, 10)
        for ri in range(2, min(len(rows) + 3, 102)):  # sample first 100 rows
            v = ws.cell(ri, ci).value
            if v:
                max_len = max(max_len, min(len(str(v)) + 2, 40))
        ws.column_dimensions[get_column_letter(ci)].width = max_len

    return ws


wb = openpyxl.Workbook()
now_str = datetime.now().strftime("%Y%m%d_%H%M%S")

with get_conn() as conn:

    # ── rebate_records ─────────────────────────────────────────────────────────
    rows = conn.execute("""
        SELECT id, batch_id, customer_account_id, program_tier_id,
               total_volume, total_commission,
               fx_lot, fx_commission,
               hang_hoa_lot, hang_hoa_commission,
               index_lot, index_commission,
               crypto_lot, crypto_commission,
               sharecfd_lot, sharecfd_commission,
               bond_lot, bond_commission,
               synthetic_lot, synthetic_commission,
               rebate_amount, status
        FROM rebate_records
        ORDER BY batch_id, customer_account_id
    """).fetchall()
    cols = [
        "id","batch_id","customer_account_id","program_tier_id",
        "total_volume","total_commission",
        "fx_lot","fx_commission","hang_hoa_lot","hang_hoa_commission",
        "index_lot","index_commission","crypto_lot","crypto_commission",
        "sharecfd_lot","sharecfd_commission","bond_lot","bond_commission",
        "synthetic_lot","synthetic_commission",
        "rebate_amount","status",
    ]
    dump_table(wb, "rebate_records", rows, cols, first=True)
    print(f"rebate_records:       {len(rows):>4} rows  |  {len(cols)} cols")

    # ── commission_records ─────────────────────────────────────────────────────
    rows = conn.execute("""
        SELECT id, batch_id, use_id, trading_account,
               total_volume, total_commission,
               fx_lot, fx_commission,
               hang_hoa_lot, hang_hoa_commission,
               index_lot, index_commission,
               crypto_lot, crypto_commission,
               sharecfd_lot, sharecfd_commission,
               bond_lot, bond_commission,
               synthetic_lot, synthetic_commission,
               lots_type, uploaded_at
        FROM commission_records
        ORDER BY batch_id, use_id
    """).fetchall()
    cols = [
        "id","batch_id","use_id","trading_account",
        "total_volume","total_commission",
        "fx_lot","fx_commission","hang_hoa_lot","hang_hoa_commission",
        "index_lot","index_commission","crypto_lot","crypto_commission",
        "sharecfd_lot","sharecfd_commission","bond_lot","bond_commission",
        "synthetic_lot","synthetic_commission",
        "lots_type","uploaded_at",
    ]
    dump_table(wb, "commission_records", rows, cols)
    print(f"commission_records:   {len(rows):>4} rows  |  {len(cols)} cols")

    # ── rebate_batches ─────────────────────────────────────────────────────────
    rows = conn.execute("""
        SELECT id, period_date, created_by, status,
               total_rows, total_amount, notes, created_at
        FROM rebate_batches
        ORDER BY period_date DESC
    """).fetchall()
    cols = ["id","period_date","created_by","status",
            "total_rows","total_amount","notes","created_at"]
    dump_table(wb, "rebate_batches", rows, cols)
    print(f"rebate_batches:       {len(rows):>4} rows  |  {len(cols)} cols")

    # ── program_tiers ──────────────────────────────────────────────────────────
    rows = conn.execute("""
        SELECT id, program_id, tier_number, label, target_lot, reward_value, reward_type
        FROM program_tiers
        ORDER BY program_id, tier_number
    """).fetchall()
    cols = ["id","program_id","tier_number","label","target_lot","reward_value","reward_type"]
    dump_table(wb, "program_tiers", rows, cols)
    print(f"program_tiers:        {len(rows):>4} rows  |  {len(cols)} cols")

    # ── programs ───────────────────────────────────────────────────────────────
    rows = conn.execute("""
        SELECT id, name, type, is_active, display_order
        FROM programs
        ORDER BY display_order
    """).fetchall()
    cols = ["id","name","type","is_active","display_order"]
    dump_table(wb, "programs", rows, cols)
    print(f"programs:             {len(rows):>4} rows  |  {len(cols)} cols")

    # ── customer_accounts ──────────────────────────────────────────────────────
    rows = conn.execute("""
        SELECT id, customer_id, broker_id, use_id, ib_number,
               broker_email, client_name, country, client_status,
               program_id, program_status, pending_program_id,
               created_at, customer_linked_at
        FROM customer_accounts
        ORDER BY id
    """).fetchall()
    cols = [
        "id","customer_id","broker_id","use_id","ib_number",
        "broker_email","client_name","country","client_status",
        "program_id","program_status","pending_program_id",
        "created_at","customer_linked_at",
    ]
    dump_table(wb, "customer_accounts", rows, cols)
    print(f"customer_accounts:    {len(rows):>4} rows  |  {len(cols)} cols")

# Save
out = os.path.join(
    os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
    f"export_db_tables_{now_str}.xlsx"
)
wb.save(out)
print(f"\nSaved: {out}")
